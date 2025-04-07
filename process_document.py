import fitz  # PyMuPDF for PDFs
from pptx import Presentation

def extract_text_from_pdf(pdf_path):
    """Extracts text from a PDF file."""
    doc = fitz.open(pdf_path)
    text = "\n".join([page.get_text() for page in doc])
    return text if text else "No text found in PDF."

def extract_text_from_ppt(ppt_path):
    """Extracts text from a PowerPoint file."""
    prs = Presentation(ppt_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text) if text else "No text found in PPT."
