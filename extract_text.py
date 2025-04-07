import fitz  # PyMuPDF for PDFs
from pptx import Presentation  # python-pptx for PPTs
import os

def extract_text_from_pdf(pdf_path):
    """Extracts text from a PDF file."""
    text = ""
    try:
        with fitz.open(pdf_path) as doc:
            for page in doc:
                text += page.get_text("text") + "\n"
    except Exception as e:
        print(f"Error extracting text from PDF: {e}")
    return text

def extract_text_from_ppt(ppt_path):
    """Extracts text from a PPT file."""
    text = ""
    try:
        presentation = Presentation(ppt_path)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        print(f"Error extracting text from PPT: {e}")
    return text

def extract_text(file_path):
    """Determines file type and extracts text accordingly."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".pptx":
        return extract_text_from_ppt(file_path)
    else:
        return "Unsupported file format"
